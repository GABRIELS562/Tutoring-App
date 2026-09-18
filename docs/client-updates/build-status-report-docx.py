"""Build the Angelo status report as a properly styled Word document.

Reads report.md (a constrained Markdown subset) and emits a .docx with a
cover page, TOC, consistent spacing, styled tables and page numbers.
"""
import re
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm, RGBColor

NAVY = RGBColor(0x1F, 0x3A, 0x5F)
ACCENT = RGBColor(0xC8, 0x10, 0x2E)
GREY_TEXT = RGBColor(0x55, 0x55, 0x55)
HEADER_FILL = "1F3A5F"
BAND_FILL = "F2F5F9"
CALLOUT_FILL = "FFF4E5"
BORDER = "BFC7D1"
BODY_FONT = "Calibri"
HEAD_FONT = "Calibri"

SRC = "report.md"
OUT = "Angelo Tutoring - Status Report - 2026-09-18.docx"


# ---------- low-level helpers ----------

def set_cell_fill(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def set_cell_margins(cell, top=60, bottom=60, left=100, right=100):
    tcPr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for side, val in (("top", top), ("bottom", bottom), ("start", left), ("end", right)):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tcPr.append(mar)


def set_table_borders(table, color=BORDER, size=4):
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        borders.append(el)
    tblPr.append(borders)


def set_repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trPr.append(el)


def keep_row_together(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:cantSplit")
    el.set(qn("w:val"), "true")
    trPr.append(el)


def set_col_widths(table, widths_cm):
    table.autofit = False
    tblPr = table._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)
    for row in table.rows:
        for idx, w in enumerate(widths_cm):
            if idx < len(row.cells):
                row.cells[idx].width = Cm(w)
    for idx, w in enumerate(widths_cm):
        if idx < len(table.columns):
            table.columns[idx].width = Cm(w)


def add_field(paragraph, instr):
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar"); fld_begin.set(qn("w:fldCharType"), "begin")
    instr_el = OxmlElement("w:instrText"); instr_el.set(qn("xml:space"), "preserve"); instr_el.text = instr
    fld_sep = OxmlElement("w:fldChar"); fld_sep.set(qn("w:fldCharType"), "separate")
    txt = OxmlElement("w:t"); txt.text = " "
    fld_end = OxmlElement("w:fldChar"); fld_end.set(qn("w:fldCharType"), "end")
    for el in (fld_begin, instr_el, fld_sep, txt, fld_end):
        run._r.append(el)
    return run


def add_inline(paragraph, text, size=None, color=None, bold=None, italic=None):
    """Add text with **bold**, *italic* and `code` inline markup."""
    tokens = re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)", text)
    for tok in tokens:
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**"):
            r = paragraph.add_run(tok[2:-2]); r.bold = True
        elif tok.startswith("*") and tok.endswith("*"):
            r = paragraph.add_run(tok[1:-1]); r.italic = True
        elif tok.startswith("`") and tok.endswith("`"):
            r = paragraph.add_run(tok[1:-1]); r.font.name = "Consolas"
        else:
            r = paragraph.add_run(tok)
        if size: r.font.size = Pt(size)
        if color: r.font.color.rgb = color
        if bold is not None and not (tok.startswith("**")): r.bold = bold
        if italic is not None: r.italic = italic


def para(doc, text="", style=None, size=None, color=None, align=None, before=0, after=6, bold=None, italic=None, keep_next=False):
    p = doc.add_paragraph(style=style) if style else doc.add_paragraph()
    if text:
        add_inline(p, text, size=size, color=color, bold=bold, italic=italic)
    pf = p.paragraph_format
    pf.space_before = Pt(before); pf.space_after = Pt(after)
    if align: p.alignment = align
    if keep_next: pf.keep_with_next = True
    return p


# ---------- document styles ----------

def setup_styles(doc):
    st = doc.styles
    normal = st["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    for name, size, before, after, color in (
        ("Heading 1", 18, 18, 8, NAVY),
        ("Heading 2", 13.5, 14, 6, NAVY),
        ("Heading 3", 11.5, 10, 4, RGBColor(0x33, 0x33, 0x33)),
    ):
        h = st[name]
        h.font.name = HEAD_FONT; h.font.size = Pt(size); h.font.bold = True
        h.font.color.rgb = color
        h.element.rPr.rFonts.set(qn("w:eastAsia"), HEAD_FONT)
        h.paragraph_format.space_before = Pt(before)
        h.paragraph_format.space_after = Pt(after)
        h.paragraph_format.keep_with_next = True

    # page setup: A4 with sensible margins
    for s in doc.sections:
        s.page_height = Cm(29.7); s.page_width = Cm(21.0)
        s.top_margin = Cm(2.2); s.bottom_margin = Cm(2.0)
        s.left_margin = Cm(2.0); s.right_margin = Cm(2.0)


def add_footer(doc):
    section = doc.sections[0]
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Angelo Tutoring Platform · Status Report · 18 September 2026 · Page ")
    r.font.size = Pt(8.5); r.font.color.rgb = GREY_TEXT
    add_field(p, "PAGE").font.size = Pt(8.5)


def cover_page(doc):
    p = para(doc, before=140, after=6)
    r = p.add_run("Angelo Tutoring Platform"); r.font.size = Pt(30); r.bold = True; r.font.color.rgb = NAVY
    p = para(doc, after=4)
    r = p.add_run("Project Status Report and Decisions Required"); r.font.size = Pt(16); r.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
    p = para(doc, after=36)
    r = p.add_run("18 September 2026"); r.font.size = Pt(12); r.font.color.rgb = GREY_TEXT

    # a slim accent rule
    t = doc.add_table(rows=1, cols=1); t.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_col_widths(t, [4.0])
    c = t.rows[0].cells[0]; set_cell_fill(c, "C8102E"); c.paragraphs[0].paragraph_format.space_after = Pt(0)
    c.paragraphs[0].add_run(" ").font.size = Pt(2)

    para(doc, after=0, before=30)
    meta = doc.add_table(rows=3, cols=2)
    set_col_widths(meta, [3.5, 10])
    for i, (k, v) in enumerate((("Prepared for", "[Client name]"), ("Prepared by", "[Your company]"), ("Version", "1.0 — for discussion"))):
        a, b = meta.rows[i].cells
        a.paragraphs[0].add_run(k).bold = True
        b.paragraphs[0].add_run(v)
        for cell in (a, b):
            cell.paragraphs[0].paragraph_format.space_after = Pt(4)
    doc.add_page_break()

    h = doc.add_paragraph("Contents", style="Heading 1")
    p = doc.add_paragraph()
    add_field(p, 'TOC \\o "1-2" \\h \\z \\u')
    p2 = para(doc, "Right-click the contents and choose “Update Field” if page numbers do not appear.", size=8.5, color=GREY_TEXT, italic=True)
    doc.add_page_break()


def enable_update_fields(doc):
    settings = doc.settings.element
    el = OxmlElement("w:updateFields"); el.set(qn("w:val"), "true")
    settings.append(el)


# ---------- tables ----------

def widths_for(headers, n_cols):
    """Column widths in cm for the 17cm text width, by table shape."""
    usable = 17.0
    key = tuple(headers)
    presets = {
        ("Area", "Complete", "Comment"): [7.5, 2.2, 7.3],
        ("Feature", "Status", "Detail"): [3.6, 2.9, 10.5],
        ("Requirement", "Service", "Rationale", "Launch cost", "Production cost"): [2.8, 2.4, 5.4, 3.2, 3.2],
        ("Active students", "Worst case per month", "Likely per month (one quarter of allowance)"): [4.5, 5.5, 7.0],
        ("Rank", "Channel", "Subscribers", "Style", "Typical length", "Assessment"): [1.1, 3.0, 2.8, 3.6, 2.0, 4.5],
        ("Rank", "Channel", "Subscribers", "Covers", "Assessment"): [1.1, 3.4, 3.0, 4.0, 5.5],
        ("Loaded", "Code", "Title", "Channel", "Views", "Length"): [1.4, 2.4, 6.6, 3.2, 1.9, 1.5],
    }
    if key in presets:
        return presets[key]
    return [usable / n_cols] * n_cols


def add_table(doc, headers, rows, aligns):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t)
    widths = widths_for(headers, len(headers))
    hdr = t.rows[0]
    set_repeat_header(hdr); keep_row_together(hdr)
    for i, h in enumerate(headers):
        c = hdr.cells[i]; set_cell_fill(c, HEADER_FILL); set_cell_margins(c)
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
        r = p.add_run(h); r.bold = True; r.font.size = Pt(9); r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        if aligns[i] == "r": p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        if aligns[i] == "c": p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for ri, row in enumerate(rows):
        r_ = t.add_row(); keep_row_together(r_)
        for i, val in enumerate(row):
            c = r_.cells[i]; set_cell_margins(c)
            if ri % 2 == 1: set_cell_fill(c, BAND_FILL)
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.1
            txt = val
            if txt == "✓":
                rr = p.add_run("✓"); rr.bold = True; rr.font.color.rgb = RGBColor(0x1E, 0x7B, 0x34); rr.font.size = Pt(10)
            else:
                add_inline(p, txt, size=9)
            if aligns[i] == "r": p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            if aligns[i] == "c": p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_col_widths(t, widths)
    para(doc, after=8)  # spacing after table
    return t


def callout(doc, text):
    t = doc.add_table(rows=1, cols=1); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t, color="E8C27A", size=6)
    c = t.rows[0].cells[0]; set_cell_fill(c, CALLOUT_FILL); set_cell_margins(c, 120, 120, 180, 180)
    p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.15
    add_inline(p, text, size=10)
    set_col_widths(t, [17.0])
    para(doc, after=8)


# ---------- markdown subset parser ----------

def parse_table(lines):
    rows = [[c.strip() for c in ln.strip().strip("|").split("|")] for ln in lines]
    headers = rows[0]
    sep = rows[1]
    aligns = []
    for s in sep:
        if s.startswith(":") and s.endswith(":"): aligns.append("c")
        elif s.endswith(":"): aligns.append("r")
        else: aligns.append("l")
    return headers, rows[2:], aligns


CALLOUT_PREFIXES = ("The application has only ever been run", "**A constraint you should know about.**")


def build(doc, md):
    lines = md.split("\n")
    i = 0
    # skip YAML front matter
    if lines[0].strip() == "---":
        i = lines.index("---", 1) + 1
    # skip the prepared-for block (rendered on the cover)
    while i < len(lines) and (lines[i].startswith("**Prepared") or lines[i].startswith("**Version") or not lines[i].strip()):
        i += 1
    list_num = 0
    while i < len(lines):
        ln = lines[i]
        if ln.startswith("# "):
            doc.add_paragraph(ln[2:].strip(), style="Heading 1"); i += 1; continue
        if ln.startswith("## "):
            doc.add_paragraph(ln[3:].strip(), style="Heading 2"); i += 1; continue
        if ln.startswith("### "):
            doc.add_paragraph(ln[4:].strip(), style="Heading 3"); i += 1; continue
        if ln.strip() == "\\newpage":
            doc.add_page_break(); i += 1; continue
        if ln.startswith("|"):
            block = []
            while i < len(lines) and lines[i].startswith("|"):
                block.append(lines[i]); i += 1
            headers, rows, aligns = parse_table(block)
            add_table(doc, headers, rows, aligns); continue
        m = re.match(r"^(\d+)\.\s+(.*)", ln)
        if m:
            # numbered item; gather continuation lines
            text = m.group(2); i += 1
            while i < len(lines) and lines[i].startswith("   ") and lines[i].strip():
                text += " " + lines[i].strip(); i += 1
            p = doc.add_paragraph()
            pf = p.paragraph_format; pf.left_indent = Cm(1.0); pf.first_line_indent = Cm(-0.75)
            pf.space_after = Pt(5); pf.line_spacing = 1.15
            r = p.add_run(f"{m.group(1)}.\t"); r.bold = True; r.font.color.rgb = NAVY
            add_inline(p, text)
            # tab stop for hanging indent
            tabs = pf.tab_stops; tabs.add_tab_stop(Cm(1.0))
            continue
        if ln.startswith("- "):
            text = ln[2:].strip(); i += 1
            while i < len(lines) and lines[i].startswith("  ") and lines[i].strip() and not lines[i].startswith("- "):
                text += " " + lines[i].strip(); i += 1
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(4); p.paragraph_format.line_spacing = 1.15
            add_inline(p, text); continue
        if not ln.strip():
            i += 1; continue
        # paragraph (may span multiple lines)
        text = ln.strip(); i += 1
        while i < len(lines) and lines[i].strip() and not re.match(r"^(#|\||- |\d+\.\s|\\newpage)", lines[i]):
            text += " " + lines[i].strip(); i += 1
        if text.startswith(CALLOUT_PREFIXES):
            callout(doc, text); continue
        if text.startswith("*Note:") and text.endswith("*"):
            para(doc, text[1:-1], size=9, color=GREY_TEXT, italic=True, after=4); continue
        para(doc, text)


def main():
    md = open(SRC, encoding="utf-8").read()
    doc = Document()
    setup_styles(doc)
    add_footer(doc)
    cover_page(doc)
    build(doc, md)
    enable_update_fields(doc)
    doc.save(OUT)
    print("wrote", OUT)


if __name__ == "__main__":
    main()
